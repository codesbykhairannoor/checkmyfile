import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Configure Section 1: Page Margins (IEEE standard: top 0.75 in, bottom 1.0 in, left 0.625 in, right 0.625 in)
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(0.625)
section.right_margin = Inches(0.625)

# Add Page Number in Footer
footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_foot = p_foot.add_run()
fld_xml = '<w:fldSimple ' + nsdecls('w') + ' w:instr="PAGE"/>'
r_foot._r.append(parse_xml(fld_xml))

# Running Header
header = section.header
p_head = header.paragraphs[0]
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_head = p_head.add_run("IEEE TRANSACTIONS ON DEPENDABLE AND SECURE COMPUTING (PREPRINT 2026)")
r_head.font.name = 'Times New Roman'
r_head.font.size = Pt(8.5)
r_head.font.color.rgb = RGBColor(128, 128, 128)

# Title & Abstract (Section 1: 1 column)
p_title = doc.add_paragraph("Zero-Egress Document Systems Using Client-Side WebAssembly")
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Test Tab-based Equation (NO TABLE MOVE HANDLE [+])
# Printable width: 8.5 - 2*0.625 = 7.25 inches
# Center: 3.625 in, Right: 7.25 in
p_eq = doc.add_paragraph()
p_eq.paragraph_format.space_before = Pt(6)
p_eq.paragraph_format.space_after = Pt(6)
tab_stops = p_eq.paragraph_format.tab_stops
tab_stops.add_tab_stop(Inches(3.625), WD_TAB_ALIGNMENT.CENTER)
tab_stops.add_tab_stop(Inches(7.25), WD_TAB_ALIGNMENT.RIGHT)

r_eq_tab = p_eq.add_run("\t")
r_eq = p_eq.add_run("T_wasm(Doc) = T_local_compute(Doc) + T_DOM_hydration")
r_eq.font.name = 'Times New Roman'
r_eq.italic = True
r_eq_end = p_eq.add_run("\t(2)")
r_eq_end.font.name = 'Times New Roman'
r_eq_end.bold = True

doc.save("test_features_out.docx")
print("test_features_out.docx successfully created!")
